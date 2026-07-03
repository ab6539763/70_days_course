#!/usr/bin/env python3
"""将 Markdown 课件批量导出为 PPT（依赖 python-pptx）"""

import re
import sys
from pathlib import Path

try:
    from pptx import Presentation
    from pptx.util import Inches, Pt
except ImportError:
    print("请安装: pip install python-pptx")
    sys.exit(1)

ROOT = Path(__file__).resolve().parent.parent
DAYS_DIR = ROOT / "days"
EXPORT_DIR = ROOT / "exports" / "ppt"


def parse_sections(md_text: str) -> list[tuple[str, str]]:
    """按 ## 标题拆分为幻灯片"""
    sections = []
    current_title = "封面"
    current_body: list[str] = []

    for line in md_text.split("\n"):
        if line.startswith("## "):
            if current_body:
                sections.append((current_title, "\n".join(current_body)))
            current_title = line[3:].strip()
            current_body = []
        else:
            current_body.append(line)

    if current_body:
        sections.append((current_title, "\n".join(current_body)))
    return sections


def add_slide(prs: Presentation, title: str, body: str):
    layout = prs.slide_layouts[1]  # Title and Content
    slide = prs.slides.add_slide(layout)
    slide.shapes.title.text = title[:100]

    tf = slide.placeholders[1].text_frame
    # 清理 markdown 标记，取前 8 行
    clean_lines = []
    for line in body.split("\n"):
        line = re.sub(r"[#*`>|]", "", line).strip()
        if line and not line.startswith("---"):
            clean_lines.append(line)
        if len(clean_lines) >= 8:
            break

    if clean_lines:
        tf.text = clean_lines[0]
        for line in clean_lines[1:]:
            p = tf.add_paragraph()
            p.text = line[:200]
            p.font.size = Pt(14)


def export_ppt():
    EXPORT_DIR.mkdir(parents=True, exist_ok=True)

    for md_path in sorted(DAYS_DIR.glob("day*.md")):
        md_text = md_path.read_text(encoding="utf-8")
        day_title = md_text.split("\n")[0].replace("# ", "")

        prs = Presentation()
        # 封面
        cover = prs.slides.add_slide(prs.slide_layouts[0])
        cover.shapes.title.text = day_title
        cover.placeholders[1].text = "零基础大模型应用开发 70 天培训课程"

        for title, body in parse_sections(md_text):
            if title in ("课程衔接说明",):
                continue
            add_slide(prs, title, body)

        out = EXPORT_DIR / f"{md_path.stem}.pptx"
        prs.save(str(out))
        print(f"✅ {out.name}")


if __name__ == "__main__":
    export_ppt()
